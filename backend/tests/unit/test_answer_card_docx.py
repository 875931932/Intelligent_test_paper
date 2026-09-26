"""答题卡 docx 导出：模板页眉页脚/页面设置保留，正文按数据重建三种可编辑作答区。"""
import io

from docx import Document as DocxDocument
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn

import app.services.answer_card_docx as acd
from app.services.answer_card_docx import export_answer_card_docx


class _FakeSession:
    """只支撑 _paper_meta 需要的 scalar()；get_paper_version 已被 monkeypatch 掉。"""

    def scalar(self, statement):
        return "大模型调优与部署技术"


def _questions() -> list[dict]:
    return [
        {"item_index": 1, "question_type": "single_choice", "stem": "题干一", "answer": "A", "score": 2,
         "options": ["甲", "乙"], "difficulty": "medium"},
        {"item_index": 2, "question_type": "true_false", "stem": "题干二", "answer": False, "score": 2,
         "options": [], "difficulty": "medium"},
        {"item_index": 3, "question_type": "fill_blank", "stem": "填空题干", "answer": "LLM", "score": 2,
         "options": [], "difficulty": "medium"},
        {"item_index": 4, "question_type": "short_answer", "stem": "简述题", "answer": "要点", "score": 5,
         "options": [], "difficulty": "medium"},
        {"item_index": 5, "question_type": "comprehensive", "stem": "综合一", "answer": "", "score": 10,
         "options": [], "difficulty": "hard",
         "subquestions": [{"prompt": "（1）补全代码", "score": 6, "answer": "x"}]},
        {"item_index": 6, "question_type": "comprehensive", "stem": "综合二", "answer": "", "score": 10,
         "options": [], "difficulty": "hard",
         "subquestions": [{"prompt": "（1）说明原因", "score": 4, "answer": "y"}]},
    ]


def _export(monkeypatch, questions: list[dict] | None = None) -> bytes:
    pv = {
        "id": "pv1",
        "exam_project_id": "p1",
        "project_name": "第一学期",
        "version_no": 6,
        "total_score": 4,
        "status": "candidate",
        "questions": _questions() if questions is None else questions,
    }
    # docx 模块按名字导入 get_paper_version，打在它自己的命名空间上
    monkeypatch.setattr(acd, "get_paper_version", lambda *a, **k: pv)
    return export_answer_card_docx(_FakeSession(), "pv1", course_id="c1")


def test_docx_keeps_template_header_footer_and_page_setup(monkeypatch):
    data = _export(monkeypatch)
    assert data[:2] == b"PK"  # docx 即 zip
    doc = DocxDocument(io.BytesIO(data))
    # 页眉（装订线 + 考场/学号栏）与页脚页码域来自模板，清正文时必须原样保留
    header_text = "".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "考场" in header_text and "学号" in header_text
    footer_text = "".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "页" in footer_text
    assert "答题卡 ｜ 试卷版本 v6" in footer_text
    # 页面设置（A4）
    section = doc.sections[0]
    assert abs(section.page_width.cm - 21.0) < 0.01
    assert abs(section.page_height.cm - 29.7) < 0.01


def test_docx_builds_regions_from_data(monkeypatch):
    doc = DocxDocument(io.BytesIO(_export(monkeypatch)))
    text = "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    )
    # 信息头 + 考生信息栏 + 题次表
    assert "答题卡" in text and "第一学期" in text
    assert "大模型调优与部署技术" in text
    assert "学号" in text and "座位号" in text
    assert "题次" in text and "总 分" in text and "评卷人" in text
    # 客观题格子表：单选、判断各一节 → 两张 2×2 格子（题号行 + 空白答案格）
    grids = [t for t in doc.tables if len(t.rows) == 2 and len(t.columns) == 2 and t.cell(0, 0).text == "题号"]
    assert len(grids) == 2
    assert [g.cell(0, 1).text for g in grids] == ["1", "2"]
    assert all(g.cell(1, 1).text == "" for g in grids)
    # 填空题一题一线（题号 + 下划线）
    fills = [p.text for p in doc.paragraphs if p.text.startswith("3. ")]
    assert fills and fills[0].startswith("3. _")
    # 5 节标题 + 得分/评卷人小格（无框 1×2 布局表，仅小格描边）
    heads = [t for t in doc.tables if len(t.rows) == 2 and len(t.columns) == 3 and t.cell(0, 0).text != "题号"]
    assert len(heads) == 5
    mark_cells = [c.text for t in heads for r in t.rows for c in r.cells]
    assert mark_cells.count("得分") == 5 and mark_cells.count("评卷人") == 5
    # 固定布局下列宽以 tblGrid 为准：节标题格 11.37cm + 得分/评卷人各 2.5cm（只写 tcW 会被 Word 忽略）
    grid_twips = [int(c.get(qn("w:w"))) for c in heads[0]._tbl.tblGrid.gridCol_lst]
    assert [round(w / 567, 2) for w in grid_twips] == [11.37, 2.5, 2.5]
    # 节标题表每个原始 w:p 都带 keepNext——竖向合并延续格不进 row.cells，
    # 漏掉它 Word 会把节标题孤立在上一页（keepNext 链断裂）
    for head in heads:
        for p_el in head._tbl.iter(qn("w:p")):
            assert p_el.get_or_add_pPr().find(qn("w:keepNext")) is not None
    # 不出题面与答案（答题卡只承接作答）
    assert "题干一" not in text and "简述题" not in text and "【答案】" not in text
    assert "补全下面" not in text and "说明原因" not in text


def test_docx_boxes_are_exact_height_and_comprehensive_pages_break(monkeypatch):
    doc = DocxDocument(io.BytesIO(_export(monkeypatch)))
    boxes = [t for t in doc.tables if len(t.rows) == 1 and len(t.columns) == 1 and not t.cell(0, 0).text.strip()]
    # 简答 1 个小框 + 综合 2 个整页大框（id 栏有文字，不算框）
    assert len(boxes) == 3
    for box in boxes:
        row = box.rows[0]
        assert row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY  # 固定高 + 整行不跨页
        tr_pr = row._tr.get_or_add_trPr()
        assert tr_pr.find(qn("w:cantSplit")) is not None
    simple = [b for b in boxes if abs(b.rows[0].height.cm - 4.4) < 0.05]
    assert len(simple) == 1
    full = [b for b in boxes if b.rows[0].height.cm > 15]
    assert len(full) == 2
    # 首题随节标题（扣掉标题高），第二题整页铺满 → 首题框更矮
    assert full[0].rows[0].height.cm < full[1].rows[0].height.cm
    # 题号行：综合第二题另起一页，简答题不换页；题号行与所属大框不拆页
    labels = {p.text: p for p in doc.paragraphs if p.text.endswith(".")}
    assert labels["6."]._p.get_or_add_pPr().find(qn("w:pageBreakBefore")) is not None
    assert labels["4."]._p.get_or_add_pPr().find(qn("w:pageBreakBefore")) is None
    assert labels["5."]._p.get_or_add_pPr().find(qn("w:keepNext")) is not None
