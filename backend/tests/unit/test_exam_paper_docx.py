"""考试卷 docx 导出：模板页眉页脚/页面设置保留，题面按 HTML 学生卷同构渲染。"""
import io

from docx import Document as DocxDocument
from docx.oxml.ns import qn

import app.services.exam_paper_docx as epd
from app.services.exam_paper_docx import export_exam_paper_docx


class _FakeSession:
    """只支撑 _paper_meta 需要的 scalar()；get_paper_version 已被 monkeypatch 掉。"""

    def scalar(self, statement):
        return "大模型调优与部署技术"


def _questions() -> list[dict]:
    return [
        {"item_index": 1, "question_type": "single_choice", "stem": "1. 题干一", "answer": "A", "score": 2,
         "options": ["甲", "乙", "丙", "丁"], "difficulty": "medium"},
        {"item_index": 2, "question_type": "true_false", "stem": "乙", "answer": True, "score": 2,
         "options": [], "difficulty": "medium"},
        {"item_index": 3, "question_type": "comprehensive", "answer": "", "score": 10, "difficulty": "hard",
         "stem": (
             "（10分）补全下面的加载代码。\n"
             "```python\nloader = WebBaseLoader(url)\nraw = ______(1)______()\n```\n"
             "特性约束如下表所示。\n"
             "| 评价指标 | TurboMind | PyTorchEngine |\n"
             "|:---|:---:|---:|\n"
             "| 初始化资源占用 | 高 | 低 |\n"
             "| 单边行 | 值 |\n"
             "请结合上表回答。"
         ),
         "options": [],
         "subquestions": [
             {"prompt": "1. 请在不改变整体结构的前提下补全代码", "score": 6, "answer": "raw = loader.load()"},
             {"prompt": "（2）可以从哪些方向优化？", "score": 4, "answer": "换 chunk_size"},
         ]},
    ]


def _export(monkeypatch, questions: list[dict] | None = None) -> bytes:
    pv = {
        "id": "pv1",
        "exam_project_id": "p1",
        "project_name": "第一学期",
        "version_no": 6,
        "total_score": 14,
        "status": "candidate",
        "questions": _questions() if questions is None else questions,
    }
    # docx 模块按名字导入 get_paper_version，打在它自己的命名空间上
    monkeypatch.setattr(epd, "get_paper_version", lambda *a, **k: pv)
    return export_exam_paper_docx(_FakeSession(), "pv1", course_id="c1")


def test_docx_keeps_template_header_footer_and_page_setup(monkeypatch):
    data = _export(monkeypatch)
    assert data[:2] == b"PK"  # docx 即 zip
    doc = DocxDocument(io.BytesIO(data))
    # 页眉（装订线 + 考场/座位号/专业名称/学号栏）与页脚页码域来自模板，清正文时必须原样保留
    header_text = "".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "考场" in header_text and "学号" in header_text and "专业名称" in header_text
    footer_text = "".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "页" in footer_text
    assert "学生卷 ｜ 试卷版本 v6" in footer_text
    section = doc.sections[0]
    assert abs(section.page_width.cm - 21.0) < 0.01
    assert abs(section.page_height.cm - 29.7) < 0.01


def test_docx_renders_stems_options_sections_like_html(monkeypatch):
    doc = DocxDocument(io.BytesIO(_export(monkeypatch)))
    paragraphs = [p.text for p in doc.paragraphs]
    text = "\n".join(paragraphs + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    # 卷面头：标题 + 副题 + 信息头 + 题次表
    assert "考试卷" in text and "第一学期" in text
    assert "大模型调优与部署技术" in text
    score_tables = [t for t in doc.tables if t.cell(0, 0).text == "题次"]
    assert len(score_tables) == 1 and len(score_tables[0].rows) == 3
    # 题干自带编号/分值已剥，题号只出现一次；题号按类型从 1（单选 1.、判断 1.、综合 1.）；
    # 客观题补作答括号（单选宽、判断窄）
    assert "1. 题干一（  ）" in paragraphs
    assert "1. 乙（ ）" in paragraphs
    assert "1. 1." not in text
    assert paragraphs.count("（  ）" ) == 0  # 括号并进题号段，不独立成段
    # 选项独立成段并缩进（≈HTML padding-left）
    assert "A. 甲" in paragraphs and "D. 丁" in paragraphs
    option_p = next(p for p in doc.paragraphs if p.text.startswith("A. "))
    assert abs(option_p.paragraph_format.left_indent.cm - 0.75) < 0.01
    # 分节标题带学生卷去向提示
    assert "一、单选题（共1题，每题2分，共2分）（将答案写在答题纸上）" in paragraphs
    assert "（将答案写在答题纸上，对的打钩 √ ，错的打叉 ×）" in text
    # 综合题：题干首段挂题号（综合题第 1 题），分问（1）…（6分）（2）…（4分）
    assert "1. 补全下面的加载代码。" in paragraphs
    assert "（1）请在不改变整体结构的前提下补全代码（6分）" in paragraphs
    assert "（2）可以从哪些方向优化？（4分）" in paragraphs
    # 学生卷不出答案与难度
    assert "【答案】" not in text and "难度" not in text
    assert "raw = loader.load()" not in text and "换 chunk_size" not in text


def test_docx_code_fence_becomes_single_cell_box(monkeypatch):
    doc = DocxDocument(io.BytesIO(_export(monkeypatch)))
    boxes = [t for t in doc.tables if len(t.rows) == 1 and len(t.columns) == 1]
    assert len(boxes) == 1  # 范本同款 1×1 描边代码框
    code_text = boxes[0].cell(0, 0).text
    assert "loader = WebBaseLoader(url)" in code_text and "raw = ______(1)______()" in code_text
    assert "```" not in code_text and "python" not in code_text  # 围栏与语言标记不入卷


def test_docx_pipe_table_becomes_shaded_real_table(monkeypatch):
    doc = DocxDocument(io.BytesIO(_export(monkeypatch)))
    md = [t for t in doc.tables if t.cell(0, 0).text == "评价指标"]
    assert len(md) == 1
    table = md[0]
    # 表头 + 2 数据行，列数以表头为准；少列行补齐（与 HTML 规则一致）
    assert len(table.rows) == 3 and len(table.columns) == 3
    assert [c.text for c in table.rows[0].cells] == ["评价指标", "TurboMind", "PyTorchEngine"]
    assert [c.text for c in table.rows[1].cells] == ["初始化资源占用", "高", "低"]
    assert [c.text for c in table.rows[2].cells] == ["单边行", "值", ""]
    # 表头浅灰底纹（HTML .md-table th 同款）
    shd = table.cell(0, 0)._tc.get_or_add_tcPr().find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "F5F5F7"
    # 表格之外的正文仍在，竖线原文不再裸露
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "特性约束如下表所示。" in text and "请结合上表回答。" in text
    assert "| 评价指标" not in text and "|---" not in text


def test_docx_does_not_double_answer_blank(monkeypatch):
    questions = [
        {"item_index": 1, "question_type": "single_choice", "stem": "甲（  ）", "answer": "A", "score": 2,
         "options": ["a", "b"], "difficulty": "medium"},
    ]
    text = "\n".join(
        p.text for p in DocxDocument(io.BytesIO(_export(monkeypatch, questions))).paragraphs
    )
    assert text.count("（  ）") == 1
