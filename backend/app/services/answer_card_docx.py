"""答题卡 Word（.docx）导出：可编辑的格子表 / 横线 / 矩形大框三种作答区。

模板 = docs/素材/答卷A卷.doc 经 Word 另存出的 .docx（入库路径见 _TEMPLATE，作为
运行时资源随代码提交，服务器无需安装 Word）。导出时打开模板并清空正文，原样保留
页眉（左侧装订线 + 考场/座位号/学号栏）、页脚（第 X 页 共 Y 页域）与页面设置；
正文按当前试卷版本数据重建，字体字号（黑体标题/新宋体横线/12pt 表格）、0.5pt 细线
与行距均对齐范本，结构与 HTML 答题卡导出同构（见 export_answer_card_html）。
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy.orm import Session

from app.services.paper_version_service import (
    _OBJECTIVE_TYPES,
    _paper_meta,
    _section_caption,
    _section_groups,
    _trim_number,
    get_paper_version,
)

# 模板入库路径（docs/素材/答卷A卷.doc → Word 另存 .docx，见模块 docstring）
_TEMPLATE = Path(__file__).resolve().parent / "templates" / "answer_card.docx"

# 范本规格：标题/信息头/节标题黑体，填空横线新宋体，表格正文 12pt 黑体；0.5pt 细线
_FONT_HEADING = "黑体"
_FONT_LINE = "新宋体"
_FONT_BODY = "宋体"
_BORDER_SZ = 4  # w:sz 单位 1/8pt → 0.5pt，与范本表格边框一致

# 版面常量（cm）：综合题整页大框按页面可用高度扣减题号行与安全余量得出
_LABEL_RESERVE = 0.8  # 题号行及其前后间距
_HEAD_RESERVE = 2.3  # 节间距(22pt) + 节标题 + 得分/评卷人小格；首题综合框必须与节标题同页放得下，否则 keepNext 链断、标题被孤立在上一页
_BOX_SLACK = 0.4  # 页底安全余量（防 EXACTLY 大框被挤到下一页）
_SIMPLE_BOX_H = 4.4  # 简答等主观题小框（HTML 版 44mm）
_GRID_PER_ROW = 10  # 客观题格子每行题数（与 HTML 版一致）
_FILL_UNDERSCORES = 31  # 填空横线：约 40% 内容宽

_ID_ROW_TEXT = (
    "学号：＿＿＿＿＿＿＿＿　　姓名：＿＿＿＿＿＿　　考场：＿＿＿＿"
    "　　座位号：＿＿＿＿\n专业名称：＿＿＿＿＿＿"  # 整行放不下，主动断行，避免「专/业名称」被拆开
)


def _set_run(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    font: str | None = None,
    color: str | None = None,
) -> None:
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if font is not None:
        run.font.name = font  # 写 w:ascii/w:hAnsi，并把 rFonts 插到 rPr 正确位置
        r_fonts = run._r.get_or_add_rPr().find(qn("w:rFonts"))
        if r_fonts is not None:
            r_fonts.set(qn("w:eastAsia"), font)  # 中文字形走 eastAsia，不依赖主题字体


def _line_exact(paragraph, pt: float) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = Pt(pt)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY  # 后设 rule：只改 lineRule 属性


def _para(
    doc,
    text: str,
    *,
    size: float,
    bold: bool = False,
    font: str | None = None,
    color: str | None = None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    line_pt: float | None = None,
    space_before: float = 0,
    space_after: float = 0,
    keep_next: bool = False,
):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    _set_run(run, size=size, bold=bold, font=font, color=color)
    fmt = paragraph.paragraph_format
    fmt.alignment = align
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.keep_with_next = keep_next
    if line_pt is not None:
        _line_exact(paragraph, line_pt)
    return paragraph


def _mini_spacer(doc, *, pt: float = 6.0, keep_next: bool = True) -> None:
    """表与表之间的空隙。空段落的行高由段落标记（w:pPr/w:rPr）字号决定，
    直接插空段落会占一整行正文，这里把标记字号压到 pt 才能得到小间隙。"""
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_with_next = keep_next
    p_pr = paragraph._p.get_or_add_pPr()
    r_pr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(pt * 2)))  # w:sz 单位 1/2pt
    r_pr.append(sz)
    p_pr.append(r_pr)


def _table_borders(table, *, sz: int = _BORDER_SZ, color: str = "auto") -> None:
    """整表六线框（范本同款）。tblPr 里 w:tblBorders 有 schema 顺序，需找锚点插入。"""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        anchor = None
        for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"):
            anchor = tbl_pr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(borders)
        else:
            tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def _cell_borders(cell, *, sz: int = _BORDER_SZ, color: str = "auto") -> None:
    """单格四边框（节标题右侧得分/评卷人小格用；标题格保持无框）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        anchor = None
        for tag in ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"):
            anchor = tc_pr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(borders)
        else:
            tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def _row_height(row, cm: float, *, rule=WD_ROW_HEIGHT_RULE.AT_LEAST) -> None:
    row.height = Cm(cm)
    row.height_rule = rule


def _cant_split(row) -> None:
    """整行不跨页：大框（EXACTLY 行高）必须整块落在一页内。"""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is not None:
        return
    el = OxmlElement("w:cantSplit")
    anchor = tr_pr.find(qn("w:trHeight"))
    if anchor is not None:
        anchor.addprevious(el)
    else:
        tr_pr.append(el)


def _set_col_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False  # w:tblLayout fixed
    # 固定布局下 Word 以 tblGrid 列宽为准（tcW 只是单元格级提示，实测被忽略），两处都要写
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths_cm):
        grid_col.set(qn("w:w"), str(int(Cm(width).twips)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def _cell_text(cell, text: str, *, size: float = 12, bold: bool = False, font: str = _FONT_HEADING) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if text:
        _set_run(paragraph.add_run(str(text)), size=size, bold=bold, font=font)


def _clear_body(doc) -> None:
    """清空范本正文，保留节属性（页眉页脚引用与页面设置都在 w:sectPr 里）。"""
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _add_title_block(doc, meta: dict) -> None:
    _para(
        doc, "答题卡", size=16, bold=True, font=_FONT_HEADING,
        align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=20, space_after=4,
    )
    _para(
        doc, str(meta.get("project_name") or ""), size=10, font=_FONT_BODY,
        color="6E6E73", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8,
    )


def _add_info_lines(doc, meta: dict) -> None:
    course = str(meta.get("course_name") or "") or "＿＿＿＿＿＿"
    _para(
        doc,
        f"课程名称：{course}　　总分：{_trim_number(meta.get('total_score'))}分"
        f"　　题量：{meta.get('question_count', 0)}题",
        size=14, bold=True, font=_FONT_HEADING, line_pt=20,
    )
    _para(
        doc,
        "考试时间：＿＿＿＿分钟　　考试形式：＿＿＿＿"
        "　　试卷类型：＿＿＿＿　　学分：＿＿＿＿",
        size=14, bold=True, font=_FONT_HEADING, line_pt=20, space_after=8,
    )


def _add_id_row(doc, width_cm: float) -> None:
    """考生信息栏（学号/姓名/考场/座位号/专业名称，带框；一行放不下，专业名称另起一行）。"""
    table = doc.add_table(rows=1, cols=1)
    _table_borders(table)
    _set_col_widths(table, [width_cm])
    _row_height(table.rows[0], 0.9)
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    _set_run(paragraph.add_run(_ID_ROW_TEXT), size=10.5, font=_FONT_BODY)


def _add_score_table(doc, groups: list[dict], width_cm: float) -> None:
    """题次表：题次 | 一 | … | 总分 | 评卷人，含分数/评分行（范本同款）。"""
    count = len(groups)
    lead_w, tail_w = 1.8, 2.0
    group_w = (width_cm - lead_w - 2 * tail_w) / max(count, 1)
    table = doc.add_table(rows=3, cols=count + 3)
    _table_borders(table)
    _set_col_widths(table, [lead_w] + [group_w] * count + [tail_w, tail_w])
    total = sum(g["total"] for g in groups)
    rows = (
        ["题次"] + [g["ordinal"] for g in groups] + ["总 分", "评卷人"],
        ["分数"] + [_trim_number(g["total"]) for g in groups] + [_trim_number(total), ""],
        ["评分"] + [""] * (count + 2),
    )
    for r, values in enumerate(rows):
        row = table.rows[r]
        _row_height(row, 0.75)
        for c, value in enumerate(values):
            _cell_text(row.cells[c], value, bold=True)


def _add_section_head(doc, group: dict, width_cm: float) -> None:
    """节标题 + 右侧得分/评卷人小格：无框 1×2 布局（标题跨两行合并），
    仅小格四边描边；全表段落 keepNext，保证节标题与首个作答区不拆页。"""
    mark_w = 2.5
    table = doc.add_table(rows=2, cols=3)
    title_cell = table.cell(0, 0).merge(table.cell(1, 0))
    _set_col_widths(table, [width_cm - 2 * mark_w, mark_w, mark_w])

    paragraph = title_cell.paragraphs[0]
    _set_run(paragraph.add_run(_section_caption(group)), size=14, bold=True, font=_FONT_HEADING)
    _line_exact(paragraph, 20)

    for r, c, text in ((0, 1, "得分"), (0, 2, "评卷人"), (1, 1, ""), (1, 2, "")):
        cell = table.cell(r, c)
        _cell_borders(cell)
        _cell_text(cell, text, bold=(r == 0))
    _row_height(table.rows[0], 0.6)
    _row_height(table.rows[1], 0.75)

    # 遍历原始 w:p 而非 row.cells：竖向合并的延续格不进 row.cells，
    # 漏掉它最后一行就不带 keepNext，Word 会在节标题与首个作答区之间断页
    for p_el in table._tbl.iter(qn("w:p")):
        p_el.get_or_add_pPr().get_or_add_keepNext()


def _add_answer_grid(doc, items: list[dict], width_cm: float) -> None:
    """客观题作答格子表：题号一行、空白答案格一行，每 10 题一块（与 HTML 版一致）。"""
    for start in range(0, len(items), _GRID_PER_ROW):
        if start:
            _mini_spacer(doc, pt=4, keep_next=False)
        chunk = items[start : start + _GRID_PER_ROW]
        cols = len(chunk) + 1
        table = doc.add_table(rows=2, cols=cols)
        _table_borders(table)
        cell_w = width_cm / cols  # 整表铺满内容宽、格子等宽（HTML 版 width:100%）
        _set_col_widths(table, [cell_w] * cols)
        rows = (
            [("题号" if start == 0 else "")] + [str(q["item_index"]) for q in chunk],
            [("答案" if start == 0 else "")] + ["" for _ in chunk],
        )
        for r, values in enumerate(rows):
            row = table.rows[r]
            _row_height(row, 0.65 if r == 0 else 0.85)
            for c, value in enumerate(values):
                _cell_text(row.cells[c], value, bold=True)


def _add_fill_lines(doc, items: list[dict]) -> None:
    """填空题一题一线：题号 + 约 40% 内容宽下划线（范本同款 12pt 新宋体）。"""
    for i, q in enumerate(items):
        _para(
            doc,
            f"{q['item_index']}. " + "_" * _FILL_UNDERSCORES,
            size=12, font=_FONT_LINE, line_pt=23,
            space_before=6 if i == 0 else 0,
        )


def _add_answer_box(
    doc,
    q: dict,
    *,
    width_cm: float,
    height_cm: float,
    gap_before: float,
    page_break: bool,
) -> None:
    """主观题矩形大框：题号行 + 1×1 描边表格（EXACTLY 行高、整行不跨页）。
    综合题整页大框；page_break=True 时题号行另起一页（对应 HTML 的逐题换页）。"""
    label = _para(
        doc, f"{q['item_index']}.", size=12, font=_FONT_LINE,
        space_before=gap_before, space_after=2, keep_next=True,
    )
    if page_break:
        label.paragraph_format.page_break_before = True
    table = doc.add_table(rows=1, cols=1)
    _table_borders(table)
    _set_col_widths(table, [width_cm])
    row = table.rows[0]
    _row_height(row, height_cm, rule=WD_ROW_HEIGHT_RULE.EXACTLY)
    _cant_split(row)


def _add_sections(doc, groups: list[dict], width_cm: float, height_cm: float) -> None:
    full_h = height_cm - _LABEL_RESERVE - _BOX_SLACK
    first_full_h = height_cm - _HEAD_RESERVE - _LABEL_RESERVE - _BOX_SLACK
    for index, group in enumerate(groups):
        if index:
            _mini_spacer(doc, pt=22)  # 节间距 ≈ HTML .section margin-bottom 26px
        _add_section_head(doc, group, width_cm)
        qt = group["type"]
        if qt in _OBJECTIVE_TYPES:
            _mini_spacer(doc, pt=8)
            _add_answer_grid(doc, group["items"], width_cm)
        elif qt == "fill_blank":
            _add_fill_lines(doc, group["items"])
        else:
            full = qt == "comprehensive"
            for i, q in enumerate(group["items"]):
                if full and i > 0:
                    gap, height, brk = 0.0, full_h, True
                elif full:
                    gap, height, brk = 8.0, first_full_h, False
                else:
                    gap, height, brk = (8.0 if i == 0 else 22.0), _SIMPLE_BOX_H, False
                _add_answer_box(
                    doc, q, width_cm=width_cm, height_cm=height,
                    gap_before=gap, page_break=brk,
                )


def _add_footer_note(doc, meta: dict) -> None:
    """页脚注记（默认页 + 奇偶页），与页码域同行程共存。"""
    note = f"答题卡 ｜ 试卷版本 v{meta.get('version_no', 1)} ｜ 请按题号作答，勿折叠污损"
    section = doc.sections[0]
    for footer in (section.footer, section.even_page_footer):
        try:
            paragraphs = footer.paragraphs
            paragraph = paragraphs[0].insert_paragraph_before() if paragraphs else footer.add_paragraph()
        except Exception:  # noqa: BLE001 — 缺页脚部件时放弃注记，不影响导出
            continue
        _set_run(paragraph.add_run(note), size=8, font=_FONT_BODY, color="86868B")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def export_answer_card_docx(
    session: Session,
    paper_version_id: str,
    *,
    course_id: str,
) -> bytes:
    """答题卡 Word 文档字节（可编辑）：页眉页脚/页面设置取自模板，正文按数据重建。"""
    if not _TEMPLATE.exists():
        raise FileNotFoundError(f"答题卡 docx 模板缺失：{_TEMPLATE}")
    pv = get_paper_version(session, paper_version_id, course_id=course_id)
    questions = pv.get("questions", [])
    groups = _section_groups(questions)
    meta = _paper_meta(session, course_id=course_id, pv=pv)
    meta["question_count"] = len(questions)

    doc = Document(str(_TEMPLATE))
    _clear_body(doc)
    section = doc.sections[0]
    width_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    height_cm = section.page_height.cm - section.top_margin.cm - section.bottom_margin.cm

    _add_title_block(doc, meta)
    _add_info_lines(doc, meta)
    _add_id_row(doc, width_cm)
    _mini_spacer(doc, pt=10)
    _add_score_table(doc, groups, width_cm)
    _mini_spacer(doc, pt=14)
    _add_sections(doc, groups, width_cm, height_cm)
    _add_footer_note(doc, meta)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
