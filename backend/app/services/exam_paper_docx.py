"""考试卷 Word（.docx）导出：可编辑的正式卷面。

模板 = docs/素材/A卷试卷_转自DOC.docx（入库路径见 _TEMPLATE，作为运行时资源随代码
提交，服务器无需安装 Word）。导出时打开模板并清空正文，原样保留页眉（左侧装订线 +
考场/座位号/专业名称/学号栏）、页脚（第 X 页 共 Y 页域）与页面设置；正文按当前试卷
版本数据重建：卷面头（标题/信息头/题次表/页脚注记）与答题卡同构（docx_kit），
题目按题型分节渲染，结构与 HTML 学生卷导出同构（export_student_paper_html）——
题干剥编号/分值前缀、客观题补作答括号、``` 围栏转 1×1 代码框（范本同款）、
GFM 管道表格转真表格（解析规则与 HTML 共用 paper_version_service._stem_blocks）、
综合题分问（1）…（6分）。考生信息在页眉（范本同款），正文不再补底部署名栏。
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from sqlalchemy.orm import Session

from app.services.docx_kit import (
    FONT_BODY,
    FONT_HEADING,
    FONT_LINE,
    cell_shade,
    cell_text,
    clear_body,
    footer_note,
    info_lines,
    line_exact,
    mini_spacer,
    para,
    score_table,
    set_col_widths,
    set_run,
    table_borders,
    title_block,
)
from app.services.paper_version_service import (
    _ANSWER_BLANK_RE,
    _OBJECTIVE_TYPES,
    _answer_blank,
    _answer_hint,
    _paper_meta,
    _section_caption,
    _section_groups,
    _stem_blocks,
    _strip_stem_noise,
    _sub_prompt,
    _sub_score,
    _trim_number,
    get_paper_version,
)

# 模板入库路径（docs/素材/A卷试卷_转自DOC.docx，见模块 docstring）
_TEMPLATE = Path(__file__).resolve().parent / "templates" / "exam_paper.docx"

# 范本版式：正文 12pt、节标题/信息头 14pt 黑体、行距 22pt（≈HTML .q-stem 1.85 行高）
_BODY_PT = 12
_LINE_PT = 22
_INDENT_CM = 0.75  # 选项与分问缩进 ≈ HTML .q-options padding-left 26px
_Q_GAP_PT = 10  # 题间空段 ≈ HTML .question margin-bottom 16px
_SECTION_GAP_PT = 18  # 节间距 ≈ HTML .section margin-bottom 26px
_CODE_PT = 10.5  # 代码框字号：新宋体等宽，长行不撑出内容宽
_CODE_LINE_PT = 15
_TABLE_PT = 11  # 题干管道表格（HTML .md-table 小于正文一号）


def _body_para(
    doc,
    text: str,
    *,
    prefix: str = "",
    indent_cm: float = 0.0,
    keep_next: bool = True,
):
    """正文段：可选加粗前缀（题号/分问编号），固定 12pt + 22pt 行距。
    段内换行折叠为空格，与 HTML 渲染（浏览器折叠空白）一致。"""
    paragraph = doc.add_paragraph()
    if prefix:
        set_run(paragraph.add_run(prefix), size=_BODY_PT, bold=True)
    set_run(paragraph.add_run(text.replace("\n", " ")), size=_BODY_PT)
    fmt = paragraph.paragraph_format
    if indent_cm:
        fmt.left_indent = Cm(indent_cm)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_with_next = keep_next
    line_exact(paragraph, _LINE_PT)
    return paragraph


def _affix_blocks(blocks: list[tuple[str, Any]], *, prefix: str, suffix: str) -> list[tuple[str, Any]]:
    """题号/作答括号挂接：前缀落在首个文本块（首块不是文本就补一个空文本块），
    尾巴并进末尾文本块（全无文本块就追加一个），代码/表格块保持原样。"""
    blocks = list(blocks)
    if prefix and not (blocks and blocks[0][0] == "text"):
        blocks.insert(0, ("text", ""))
    if suffix:
        for i in range(len(blocks) - 1, -1, -1):
            if blocks[i][0] == "text":
                blocks[i] = ("text", blocks[i][1] + suffix)
                break
        else:
            blocks.append(("text", suffix))
    return blocks


def _add_code_box(doc, code: str, width_cm: float) -> None:
    """``` 围栏代码块 → 1×1 描边表格（范本同款），逐行 w:br 保住缩进。"""
    table = doc.add_table(rows=1, cols=1)
    table_borders(table)
    set_col_widths(table, [width_cm])
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(1)
    fmt.space_after = Pt(1)
    fmt.keep_with_next = True  # 与后续分问/正文不拆页（整题空段负责断链）
    line_exact(paragraph, _CODE_LINE_PT)
    lines = (code or "").split("\n")
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_run(run, size=_CODE_PT, font=FONT_LINE)
        if i < len(lines) - 1:
            run.add_break()


def _add_md_table(doc, payload: tuple[list[str], list[list[str]]], width_cm: float) -> None:
    """GFM 管道表格 → 有框真表格：表头浅灰加粗、全格居中（HTML .md-table 同款）。"""
    header, rows = payload
    cols = max(len(header), 1)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table_borders(table)
    set_col_widths(table, [width_cm / cols] * cols)
    for r, cells in enumerate([header] + rows):
        padded = (list(cells) + [""] * cols)[:cols]
        row = table.rows[r]
        for c, value in enumerate(padded):
            cell = row.cells[c]
            cell_text(cell, str(value), size=_TABLE_PT, bold=(r == 0), font=FONT_BODY)
            if r == 0:
                cell_shade(cell, "F5F5F7")
    # 全表段落 keepNext：题干表格与后续正文不拆页（与答题卡节标题表同手法）
    for p_el in table._tbl.iter(qn("w:p")):
        p_el.get_or_add_pPr().get_or_add_keepNext()


def _add_blocks(
    doc,
    blocks: list[tuple[str, Any]],
    *,
    width_cm: float,
    prefix: str = "",
    suffix: str = "",
    indent_cm: float = 0.0,
) -> None:
    """解析块 → Word：text → 正文段（首段挂加粗前缀），code → 代码框，table → 真表格。"""
    blocks = _affix_blocks(blocks, prefix=prefix, suffix=suffix)
    first_text = True
    for kind, payload in blocks:
        if kind == "text":
            _body_para(
                doc, payload,
                prefix=prefix if first_text else "",
                indent_cm=indent_cm,
            )
            first_text = False
        elif kind == "code":
            _add_code_box(doc, payload, width_cm)
        else:
            _add_md_table(doc, payload, width_cm)


def _add_question(doc, q: dict, *, width_cm: float) -> None:
    """单题（学生卷）：题号+题干+作答括号 / 选项缩进 / 分问（1）…（6分）；
    不出答案与难度（属答卷与内部信息）。整题段落 keepNext，
    题间空段（keep_next=False）收口，整题尽量不跨页。"""
    stem = _strip_stem_noise(str(q.get("stem", "")))
    qtype = q.get("question_type") or ""
    suffix = _answer_blank(qtype) if qtype in _OBJECTIVE_TYPES and not _ANSWER_BLANK_RE.search(stem) else ""
    _add_blocks(
        doc, _stem_blocks(stem), width_cm=width_cm,
        prefix=f"{q.get('item_index', 0)}. ", suffix=suffix,
    )

    for i, opt in enumerate(q.get("options") or []):
        text = opt.get("text", str(opt)) if isinstance(opt, dict) else str(opt)
        _body_para(doc, f"{chr(65 + i)}. {text}", indent_cm=_INDENT_CM)

    for i, sub in enumerate(q.get("subquestions") or [], start=1):
        score = _sub_score(sub)
        tail = f"（{_trim_number(score)}分）" if score is not None else ""
        _add_blocks(
            doc, _stem_blocks(_sub_prompt(sub)), width_cm=width_cm,
            prefix=f"（{i}）", suffix=tail, indent_cm=_INDENT_CM,
        )

    mini_spacer(doc, pt=_Q_GAP_PT, keep_next=False)  # 题间空段，同时断开整题 keepNext 链


def _add_sections(doc, groups: list[dict], *, width_cm: float) -> None:
    for index, group in enumerate(groups):
        if index:
            mini_spacer(doc, pt=_SECTION_GAP_PT)  # 节间距；keep_next 默认 True，粘住节标题
        hint = _answer_hint(group["type"], with_answer=False)
        para(
            doc, _section_caption(group, hint=hint),
            size=14, bold=True, font=FONT_HEADING,
            line_pt=20, space_after=4, keep_next=True,
        )
        for q in group["items"]:
            _add_question(doc, q, width_cm=width_cm)


def export_exam_paper_docx(
    session: Session,
    paper_version_id: str,
    *,
    course_id: str,
) -> bytes:
    """考试卷 Word 文档字节（可编辑）：页眉页脚/页面设置取自模板，正文按数据重建。"""
    if not _TEMPLATE.exists():
        raise FileNotFoundError(f"考试卷 docx 模板缺失：{_TEMPLATE}")
    pv = get_paper_version(session, paper_version_id, course_id=course_id)
    questions = pv.get("questions", [])
    groups = _section_groups(questions)
    meta = _paper_meta(session, course_id=course_id, pv=pv)
    meta["question_count"] = len(questions)

    doc = Document(str(_TEMPLATE))
    clear_body(doc)
    section = doc.sections[0]
    width_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm

    title_block(doc, meta, "考试卷")
    info_lines(doc, meta)
    score_table(doc, groups, width_cm)
    mini_spacer(doc, pt=14)
    _add_sections(doc, groups, width_cm=width_cm)
    footer_note(doc, f"学生卷 ｜ 试卷版本 v{meta.get('version_no', 1)} ｜ 请将答案作答在答题卡上")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
