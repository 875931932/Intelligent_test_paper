"""Word（.docx）导出公共件：答题卡与考试卷共用的模板清理、卷面头与表格/段落工具。

模板打开后清空正文、保留 sectPr（页眉页脚与页面设置都在其中），正文按数据重建；
字体字号、0.5pt 细线与行距均对齐 docs/素材 下的命题范本。两份导出各自的作答区/
题面渲染留在各自模块，此处只放真正同构的部分（范本同款标题/信息头/题次表/页脚注记
与通用 XML 工具）。
"""
from __future__ import annotations

from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.services.paper_version_service import _trim_number

# 范本规格：标题/信息头/节标题黑体，填空横线新宋体，表格正文 12pt 黑体；0.5pt 细线
FONT_HEADING = "黑体"
FONT_LINE = "新宋体"
FONT_BODY = "宋体"
BORDER_SZ = 4  # w:sz 单位 1/8pt → 0.5pt，与范本表格边框一致


def set_run(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    font: str | None = None,
    color: str | None = None,
) -> None:
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if font is not None:
        run.font.name = font  # 写 w:ascii/w:hAnsi，并把 rFonts 插到 rPr 正确位置
        r_fonts = run._r.get_or_add_rPr().find(qn("w:rFonts"))
        if r_fonts is not None:
            r_fonts.set(qn("w:eastAsia"), font)  # 中文字形走 eastAsia，不依赖主题字体


def line_exact(paragraph, pt: float) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = Pt(pt)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY  # 后设 rule：只改 lineRule 属性


def para(
    doc,
    text: str,
    *,
    size: float,
    bold: bool = False,
    font: str | None = None,
    color: str | None = None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    line_pt: float | None = None,
    space_before: float = 0,
    space_after: float = 0,
    keep_next: bool = False,
):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_run(run, size=size, bold=bold, font=font, color=color)
    fmt = paragraph.paragraph_format
    fmt.alignment = align
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.keep_with_next = keep_next
    if line_pt is not None:
        line_exact(paragraph, line_pt)
    return paragraph


def mini_spacer(doc, *, pt: float = 6.0, keep_next: bool = True) -> None:
    """表与表之间的空隙。空段落的行高由段落标记（w:pPr/w:rPr）字号决定，
    直接插空段落会占一整行正文，这里把标记字号压到 pt 才能得到小间隙。"""
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_with_next = keep_next
    p_pr = paragraph._p.get_or_add_pPr()
    r_pr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(pt * 2)))  # w:sz 单位 1/2pt
    r_pr.append(sz)
    p_pr.append(r_pr)


def table_borders(table, *, sz: int = BORDER_SZ, color: str = "auto") -> None:
    """整表六线框（范本同款）。tblPr 里 w:tblBorders 有 schema 顺序，需找锚点插入。"""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        anchor = None
        for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"):
            anchor = tbl_pr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(borders)
        else:
            tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def cell_borders(cell, *, sz: int = BORDER_SZ, color: str = "auto") -> None:
    """单格四边框（节标题右侧得分/评卷人小格用；标题格保持无框）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        anchor = None
        for tag in ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"):
            anchor = tc_pr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(borders)
        else:
            tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def cell_shade(cell, fill: str) -> None:
    """单元格底纹（题干管道表格表头浅灰，对齐 HTML .md-table th）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        anchor = None
        for tag in ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"):
            anchor = tc_pr.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(shd)
        else:
            tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def row_height(row, cm: float, *, rule=WD_ROW_HEIGHT_RULE.AT_LEAST) -> None:
    row.height = Cm(cm)
    row.height_rule = rule


def cant_split(row) -> None:
    """整行不跨页：大框（EXACTLY 行高）必须整块落在一页内。"""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is not None:
        return
    el = OxmlElement("w:cantSplit")
    anchor = tr_pr.find(qn("w:trHeight"))
    if anchor is not None:
        anchor.addprevious(el)
    else:
        tr_pr.append(el)


def set_col_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False  # w:tblLayout fixed
    # 固定布局下 Word 以 tblGrid 列宽为准（tcW 只是单元格级提示，实测被忽略），两处都要写
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths_cm):
        grid_col.set(qn("w:w"), str(int(Cm(width).twips)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def cell_text(cell, text: str, *, size: float = 12, bold: bool = False, font: str = FONT_HEADING) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if text:
        set_run(paragraph.add_run(str(text)), size=size, bold=bold, font=font)


def clear_body(doc) -> None:
    """清空范本正文，保留节属性（页眉页脚引用与页面设置都在 w:sectPr 里）。"""
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def title_block(doc, meta: dict, title: str) -> None:
    para(
        doc, title, size=16, bold=True, font=FONT_HEADING,
        align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=20, space_after=4,
    )
    para(
        doc, str(meta.get("project_name") or ""), size=10, font=FONT_BODY,
        color="6E6E73", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8,
    )


def info_lines(doc, meta: dict) -> None:
    course = str(meta.get("course_name") or "") or "＿＿＿＿＿＿"
    para(
        doc,
        f"课程名称：{course}　　总分：{_trim_number(meta.get('total_score'))}分"
        f"　　题量：{meta.get('question_count', 0)}题",
        size=14, bold=True, font=FONT_HEADING, line_pt=20,
    )
    para(
        doc,
        "考试时间：＿＿＿＿分钟　　考试形式：＿＿＿＿"
        "　　试卷类型：＿＿＿＿　　学分：＿＿＿＿",
        size=14, bold=True, font=FONT_HEADING, line_pt=20, space_after=8,
    )


def score_table(doc, groups: list[dict], width_cm: float) -> None:
    """题次表：题次 | 一 | … | 总分 | 评卷人，含分数/评分行（范本同款）。"""
    count = len(groups)
    lead_w, tail_w = 1.8, 2.0
    group_w = (width_cm - lead_w - 2 * tail_w) / max(count, 1)
    table = doc.add_table(rows=3, cols=count + 3)
    table_borders(table)
    set_col_widths(table, [lead_w] + [group_w] * count + [tail_w, tail_w])
    total = sum(g["total"] for g in groups)
    rows = (
        ["题次"] + [g["ordinal"] for g in groups] + ["总 分", "评卷人"],
        ["分数"] + [_trim_number(g["total"]) for g in groups] + [_trim_number(total), ""],
        ["评分"] + [""] * (count + 2),
    )
    for r, values in enumerate(rows):
        row = table.rows[r]
        row_height(row, 0.75)
        for c, value in enumerate(values):
            cell_text(row.cells[c], value, bold=True)


def footer_note(doc, note: str) -> None:
    """页脚注记（默认页 + 奇偶页），与页码域同行程共存。"""
    section = doc.sections[0]
    for footer in (section.footer, section.even_page_footer):
        try:
            paragraphs = footer.paragraphs
            paragraph = paragraphs[0].insert_paragraph_before() if paragraphs else footer.add_paragraph()
        except Exception:  # noqa: BLE001 — 缺页脚部件时放弃注记，不影响导出
            continue
        set_run(paragraph.add_run(note), size=8, font=FONT_BODY, color="86868B")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
